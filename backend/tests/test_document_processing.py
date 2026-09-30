import pytest

from app.services.document_processing import (
    Block,
    DocumentValidationError,
    chunk_blocks,
    clean_text,
    parse_document,
    safe_filename,
    validate_upload,
)


def test_validate_rejects_unknown_extension():
    with pytest.raises(DocumentValidationError):
        validate_upload("malware.exe", b"MZ...", 1024)


def test_validate_rejects_spoofed_pdf():
    with pytest.raises(DocumentValidationError, match="not a valid PDF"):
        validate_upload("policy.pdf", b"hello world", 1024)


def test_validate_rejects_oversized_and_empty():
    with pytest.raises(DocumentValidationError, match="limit"):
        validate_upload("a.txt", b"x" * 2048, 1024)
    with pytest.raises(DocumentValidationError, match="empty"):
        validate_upload("a.txt", b"", 1024)


def test_validate_rejects_binary_text():
    with pytest.raises(DocumentValidationError):
        validate_upload("a.txt", b"abc\x00def", 1024)


def test_safe_filename_strips_paths():
    assert safe_filename("../../etc/passwd") == "passwd"
    assert "/" not in safe_filename("a/b\\c<>.md")


def test_markdown_sections_are_tracked():
    blocks = parse_document(b"# Refunds\n\nRefunds take 5 days.\n\n## Exchanges\n\nExchanges depend on stock.", "md")
    assert [(b.section, b.text) for b in blocks] == [
        ("Refunds", "Refunds take 5 days."),
        ("Exchanges", "Exchanges depend on stock."),
    ]


def test_clean_text_unwraps_soft_breaks():
    assert clean_text("This is a\nwrapped line.\n\nNext") == "This is a wrapped line.\n\nNext"


def test_chunking_respects_size_and_overlap():
    sentences = " ".join(f"Sentence number {i} has some words." for i in range(60))
    chunks = chunk_blocks([Block(sentences, "S")], chunk_size=300, overlap=80)
    assert len(chunks) > 3
    assert all(len(c.text) <= 300 for c in chunks)
    # consecutive chunks share overlapping text
    assert any(c1.text[-40:].split(".")[-2] in c2.text for c1, c2 in zip(chunks, chunks[1:]))
    assert [c.index for c in chunks] == list(range(len(chunks)))


def test_chunking_does_not_cross_sections():
    chunks = chunk_blocks([Block("Alpha text.", "A"), Block("Beta text.", "B")], chunk_size=1000, overlap=100)
    assert [(c.section, c.text) for c in chunks] == [("A", "Alpha text."), ("B", "Beta text.")]


def test_docx_parsing_uses_headings_as_sections():
    import io

    import docx

    d = docx.Document()
    d.add_heading("Warranty", level=1)
    d.add_paragraph("Phones carry a 1-year warranty.")
    buf = io.BytesIO()
    d.save(buf)
    data = buf.getvalue()
    assert validate_upload("w.docx", data, 10_000_000) == "docx"
    blocks = parse_document(data, "docx")
    assert [(b.section, b.text) for b in blocks] == [("Warranty", "Phones carry a 1-year warranty.")]
