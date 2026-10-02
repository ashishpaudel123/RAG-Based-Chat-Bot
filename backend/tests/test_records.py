"""Gemini batch import and record linting."""

from app.services.records import check_folder, check_record, read_text, split_batch

GOOD = """---
id: citizenship.duplicate.lost.001
title: "हराएको नागरिकताको प्रतिलिपि (Duplicate of a lost certificate)"
source_type: regulation
section: "नियम ९"
status: current
source_url: https://example.gov.np/regulation.pdf
verification_status: pending
user_questions:
  - "nagarikta harayo"
---
# हराएको नागरिकताको प्रतिलिपि (Duplicate of a lost certificate)

## सारांश (Summary)
""" + ("Body text from the regulation. " * 12) + """

## स्रोत (Sources)
Nepal Citizenship Regulation, 2063, Rule 9.
"""

BATCH = f"""Here are the records for batch 4.

FILE: citizenship.duplicate.lost.001.md
```markdown
{GOOD}```

FILE: `citizenship.bad.002.md`
```markdown
---
id: citizenship.bad.002
title: "<नेपाली शीर्षक> (<English title>)"
source_type: blogpost
status: current
verification_status: pending
---
short
```

Topics not covered: none.
"""


def test_split_batch_from_text():
    records = split_batch(BATCH)
    assert [r[0] for r in records] == ["citizenship.duplicate.lost.001", "citizenship.bad.002"]
    assert records[0][1].startswith("---\nid: citizenship.duplicate.lost.001")
    assert "```" not in records[0][1] and "Topics not covered" not in records[1][1]


def test_split_batch_from_docx_without_fences(tmp_path):
    import docx

    d = docx.Document()
    for line in ("Intro", "FILE: citizenship.duplicate.lost.001.md", *GOOD.splitlines(), "Gemini notes after"):
        d.add_paragraph(line)
    path = tmp_path / "batch.docx"
    d.save(path)
    records = split_batch(read_text(path))
    assert len(records) == 1 and check_record("citizenship.duplicate.lost.001.md", records[0][1]).ok


def test_check_record_reports_problems():
    assert check_record("citizenship.duplicate.lost.001.md", GOOD).ok
    bad = check_record("citizenship.bad.002.md", split_batch(BATCH)[1][1])
    joined = " ".join(bad.errors)
    assert "source_type" in joined and "placeholders" in joined and "too short" in joined


def test_check_folder_finds_duplicate_ids(tmp_path):
    (tmp_path / "a.md").write_text(GOOD, encoding="utf-8")
    (tmp_path / "b.md").write_text(GOOD, encoding="utf-8")
    (tmp_path / "_TEMPLATE.md").write_text("ignored", encoding="utf-8")
    reports = check_folder(tmp_path)
    assert len(reports) == 2 and any("duplicate id" in e for r in reports for e in r.errors)


def test_template_is_flagged_as_unfilled():
    from pathlib import Path

    template = Path(__file__).resolve().parents[2] / "knowledge_base" / "_TEMPLATE.md"
    assert not check_record("x.md", template.read_text(encoding="utf-8")).ok
