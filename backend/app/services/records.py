"""Knowledge-record tooling: split a Gemini batch output into record files and
lint records before ingestion.

Gemini (prompt 3 in the README workflow) answers with blocks like::

    FILE: citizenship.duplicate.lost.001.md
    ```markdown
    ---
    id: citizenship.duplicate.lost.001
    ...
    ---
    # body
    ```

The text may be pasted into a .txt/.md file or exported as a .docx (where the
code fences are often lost); both are handled.
"""

import io
import re
from dataclasses import dataclass, field
from pathlib import Path

from app.services.document_metadata import normalize_metadata, split_front_matter
from app.services.document_processing import DocumentValidationError

_FILE_LINE = re.compile(r"^\s*(?:\*\*)?FILE\s*:\s*`?([\w.\-]+?)(?:\.md)?`?(?:\*\*)?\s*$", re.I | re.M)
_FENCE = re.compile(r"^\s*```[\w-]*\s*$")
_PLACEHOLDER = re.compile(r"<(?:[^<>\n]{0,40}(?:नेपाली|English|roman|exact|official|date|category|topic|…|\.\.\.)[^<>\n]{0,60})>",
                          re.I)
_ID = re.compile(r"^[a-z0-9][a-z0-9._-]{2,119}$")

REQUIRED = ("id", "title", "source_type", "status", "verification_status")
EXPECTED_SECTIONS = ("(Summary)", "(Sources)")


def read_text(path: Path) -> str:
    """Text of a .txt/.md file, or the paragraphs of a .docx export."""
    if path.suffix.lower() == ".docx":
        import docx

        document = docx.Document(io.BytesIO(path.read_bytes()))
        return "\n".join(p.text for p in document.paragraphs)
    return path.read_text(encoding="utf-8-sig")


def split_batch(text: str) -> list[tuple[str, str]]:
    """Return [(record_id, markdown)] for every "FILE:" block in a Gemini answer."""
    matches = list(_FILE_LINE.finditer(text))
    records = []
    for i, m in enumerate(matches):
        chunk = text[m.end(): matches[i + 1].start() if i + 1 < len(matches) else len(text)]
        lines = chunk.strip("\n").splitlines()
        # drop the opening fence (and anything before the front matter)
        while lines and not lines[0].strip().startswith("---"):
            lines.pop(0)
        # drop the closing fence and any trailing commentary after it
        for j in range(len(lines) - 1, 0, -1):
            if _FENCE.match(lines[j]):
                lines = lines[:j]
                break
        body = "\n".join(lines).strip()
        if body.startswith("---"):
            records.append((m.group(1), body + "\n"))
    return records


@dataclass
class RecordReport:
    file: str
    record_id: str | None = None
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.errors


def check_record(name: str, text: str) -> RecordReport:
    report = RecordReport(file=name)
    try:
        meta, body = split_front_matter(text)
    except DocumentValidationError as exc:
        report.errors.append(str(exc))
        return report
    if not meta:
        report.errors.append("missing YAML front matter (--- ... --- at the top)")
        return report
    report.record_id = str(meta.get("id") or "")
    for key in REQUIRED:
        if meta.get(key) in (None, ""):
            report.errors.append(f"missing required field '{key}'")
    try:
        values, _ = normalize_metadata(meta)
    except DocumentValidationError as exc:
        report.errors.append(str(exc))
        values = {}
    rid = report.record_id
    if rid and not _ID.match(rid):
        report.errors.append(f"id '{rid}' should use lowercase letters, digits, dots, dashes or underscores")
    stem = Path(name).stem
    if rid and stem != rid:
        report.warnings.append(f"file name '{stem}' differs from id '{rid}'")
    if not meta.get("source_url") and values.get("source_type") not in (None, "faq", "secondary"):
        report.warnings.append("no source_url: every legal record should link to its official source")
    if not meta.get("section"):
        report.warnings.append("no section: cite the exact दफा/नियम/Article")
    if values.get("verification_status") == "verified" and not meta.get("last_verified"):
        report.warnings.append("marked verified but last_verified is empty")
    if meta.get("district") and values.get("source_type") not in ("dao_charter", "local_notice", "official_portal"):
        report.warnings.append("district is set but source_type is not local (dao_charter/local_notice)")
    if len(body.strip()) < 200:
        report.errors.append("body is too short to be useful (< 200 characters)")
    for section in EXPECTED_SECTIONS:
        if section not in body:
            report.warnings.append(f"missing section heading containing '{section}'")
    leftovers = sorted(set(_PLACEHOLDER.findall(text)))
    if leftovers:
        report.errors.append("template placeholders left in: " + ", ".join(leftovers[:5]))
    if not meta.get("user_questions"):
        report.warnings.append("no user_questions: add Nepali/Roman/English phrasings for better retrieval")
    return report


def check_folder(folder: Path) -> list[RecordReport]:
    reports, seen = [], {}
    for path in sorted(folder.rglob("*.md")):
        if path.name.lower().startswith(("readme", "_")):
            continue
        rep = check_record(str(path.relative_to(folder)), path.read_text(encoding="utf-8-sig"))
        if rep.record_id:
            if rep.record_id in seen:
                rep.errors.append(f"duplicate id (also in {seen[rep.record_id]}); use a new id or it becomes a new version")
            seen.setdefault(rep.record_id, rep.file)
        reports.append(rep)
    return reports
