"""Offline knowledge pipeline stages: validate -> parse -> clean -> chunk (+metadata).

Proposal §3.3/§3.4. Files are validated by extension, size and magic bytes
before parsing (§3.11).
"""

import io
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path

ALLOWED_TYPES = {".pdf": "pdf", ".docx": "docx", ".txt": "txt", ".md": "md"}


class DocumentValidationError(ValueError):
    pass


@dataclass
class Block:
    """A paragraph-level unit of text with its location metadata."""

    text: str
    section: str | None = None
    page: int | None = None


@dataclass
class TextChunk:
    index: int
    text: str
    section: str | None
    page: int | None


# ------------------------------------------------------------ validate ----
def validate_upload(filename: str, data: bytes, max_bytes: int) -> str:
    """Return the normalised content type or raise DocumentValidationError."""
    ext = Path(filename or "").suffix.lower()
    if ext not in ALLOWED_TYPES:
        raise DocumentValidationError(
            f"Unsupported file type '{ext or 'none'}'. Allowed: {', '.join(sorted(ALLOWED_TYPES))}"
        )
    if not data:
        raise DocumentValidationError("The uploaded file is empty")
    if len(data) > max_bytes:
        raise DocumentValidationError(f"File exceeds the {max_bytes // (1024 * 1024)} MB limit")
    kind = ALLOWED_TYPES[ext]
    if kind == "pdf" and not data.startswith(b"%PDF-"):
        raise DocumentValidationError("File content is not a valid PDF")
    if kind == "docx" and not data.startswith(b"PK\x03\x04"):
        raise DocumentValidationError("File content is not a valid DOCX document")
    if kind in ("txt", "md"):
        if b"\x00" in data[:4096]:
            raise DocumentValidationError("Text files must not contain binary data")
        try:
            data.decode("utf-8")
        except UnicodeDecodeError:
            raise DocumentValidationError("Text files must be UTF-8 encoded")
    return kind


def safe_filename(filename: str) -> str:
    name = Path(filename or "document").name
    name = re.sub(r"[^A-Za-z0-9._ -]", "_", name).strip(" .") or "document"
    return name[:200]


def title_from_filename(filename: str) -> str:
    stem = Path(filename).stem
    return re.sub(r"[_-]+", " ", stem).strip().title() or "Untitled document"


# --------------------------------------------------------------- parse ----
def parse_document(data: bytes, kind: str) -> list[Block]:
    if kind == "pdf":
        blocks = _parse_pdf(data)
    elif kind == "docx":
        blocks = _parse_docx(data)
    elif kind == "md":
        blocks = _parse_markdown(data.decode("utf-8"))
    else:
        blocks = _parse_text(data.decode("utf-8"))
    blocks = [Block(clean_text(b.text), b.section, b.page) for b in blocks]
    blocks = [b for b in blocks if b.text]
    if not blocks:
        raise DocumentValidationError("No readable text could be extracted from the document")
    return blocks


def _paragraphs(text: str) -> list[str]:
    return [p for p in re.split(r"\n\s*\n", text) if p.strip()]


def _parse_pdf(data: bytes) -> list[Block]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise DocumentValidationError("Encrypted PDFs are not supported")
        blocks: list[Block] = []
        for page_no, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            blocks.extend(Block(p, None, page_no) for p in _paragraphs(text))
        return blocks
    except DocumentValidationError:
        raise
    except Exception as exc:
        raise DocumentValidationError(f"Could not read PDF: {type(exc).__name__}") from exc


def _parse_docx(data: bytes) -> list[Block]:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentValidationError(f"Could not read DOCX: {type(exc).__name__}") from exc
    blocks: list[Block] = []
    section = None
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style.startswith("heading") or style == "title":
            section = text[:255]
            continue
        blocks.append(Block(text, section))
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                blocks.append(Block(" | ".join(dict.fromkeys(cells)), section))
    return blocks


def _parse_markdown(text: str) -> list[Block]:
    blocks: list[Block] = []
    section = None
    buf: list[str] = []

    def flush():
        if buf:
            blocks.extend(Block(p, section) for p in _paragraphs("\n".join(buf)))
            buf.clear()

    for line in text.splitlines():
        heading = re.match(r"^\s{0,3}#{1,6}\s+(.*)$", line)
        if heading:
            flush()
            section = heading.group(1).strip().strip("#").strip()[:255] or section
        else:
            buf.append(line)
    flush()
    return blocks


def _parse_text(text: str) -> list[Block]:
    return [Block(p) for p in _paragraphs(text)]


# --------------------------------------------------------------- clean ----
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def clean_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    text = _CONTROL.sub(" ", text)
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)          # re-join hyphenated line breaks
    text = re.sub(r"[ \t]*\n[ \t]*", "\n", text)
    text = re.sub(r"(?<![.:!?\n])\n(?!\n|[-*•]|\d+[.)])", " ", text)  # unwrap soft line breaks
    text = re.sub(r"[ \t]{2,}", " ", text)
    return text.strip()


# --------------------------------------------------------------- chunk ----
def _split_long(text: str, size: int) -> list[str]:
    if len(text) <= size:
        return [text]
    pieces: list[str] = []
    current = ""
    for sentence in re.split(r"(?<=[.!?])\s+", text):
        while len(sentence) > size:  # a single run-on "sentence"
            cut = sentence.rfind(" ", 0, size)
            cut = cut if cut > size // 2 else size
            if current:
                pieces.append(current)
                current = ""
            pieces.append(sentence[:cut].strip())
            sentence = sentence[cut:].strip()
        if current and len(current) + 1 + len(sentence) > size:
            pieces.append(current)
            current = sentence
        else:
            current = f"{current} {sentence}".strip()
    if current:
        pieces.append(current)
    return pieces


def chunk_blocks(blocks: list[Block], chunk_size: int, overlap: int) -> list[TextChunk]:
    """Greedy paragraph-aware chunking with character overlap.

    Paragraphs are packed into chunks up to ``chunk_size`` characters without
    crossing section boundaries; the tail of each chunk (up to ``overlap``
    characters, on a sentence boundary where possible) seeds the next chunk.
    """
    if chunk_size <= 0:
        raise ValueError("chunk_size must be positive")
    overlap = max(0, min(overlap, chunk_size // 2))

    units: list[Block] = []
    for b in blocks:
        units.extend(Block(piece, b.section, b.page) for piece in _split_long(b.text, chunk_size))

    chunks: list[TextChunk] = []
    current: list[str] = []
    cur_section: str | None = None
    cur_page: int | None = None

    def emit():
        text = "\n\n".join(current).strip()
        if text:
            chunks.append(TextChunk(len(chunks), text, cur_section, cur_page))

    def tail(text: str) -> str:
        if overlap == 0 or len(text) <= overlap:
            return "" if overlap == 0 else text
        snippet = text[-overlap:]
        boundary = re.search(r"(?<=[.!?])\s+", snippet)
        return snippet[boundary.end():] if boundary else snippet[snippet.find(" ") + 1:]

    for unit in units:
        new_section = unit.section != cur_section and current
        projected = len("\n\n".join(current + [unit.text]))
        if current and (new_section or projected > chunk_size):
            emit()
            carry = "" if new_section else tail("\n\n".join(current))
            current = [carry] if carry and len(carry) + len(unit.text) + 2 <= chunk_size else []
        if not current:
            cur_section, cur_page = unit.section, unit.page
        current.append(unit.text)
    emit()
    return chunks
